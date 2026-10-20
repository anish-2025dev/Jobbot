import copy
import io
import json
from pathlib import Path

import httpx
import pytest

from app.llm.client import LLM
from app.resume.cover import standard_answers, template_cover_letter
from app.resume.engine import _drop_one, build_recipe, prepare_jd
from app.resume.model import build_view, date_range, display_url
from app.resume.recipe import RBullet, RSection, RSkillGroup
from app.resume.render import latex_escape, to_docx, to_latex, to_text, to_typst, typst_string
from app.resume.select import estimate_lines, select
from app.resume.tailor import tailor
from app.resume.validator import check_bullet, validate
from app.vault.facts import extract_numbers
from app.vault.schema import Vault

ROOT = Path(__file__).parents[2]
EXAMPLE = json.loads((ROOT / "examples/vault.example.json").read_text())
JD_TEXT = (ROOT / "examples/jd.example.txt").read_text()
DS_JD = """Data Engineer (Fresher). Build ELT pipelines with Airflow and dbt on PostgreSQL,
model data into star schemas, publish Power BI dashboards and track experiments in MLflow.
Requirements: SQL, Python, dbt, Airflow."""


@pytest.fixture
def v():
    return Vault.model_validate(copy.deepcopy(EXAMPLE))


def _jd(v, text=JD_TEXT, title="Software Development Engineer I", company="Northwind Payments"):
    return prepare_jd(v, text, title, company)


# --- selection ------------------------------------------------------------------
def test_select_is_relevant_and_deterministic(v):
    a = select(v, _jd(v), vault_revision=3, bucket="high")
    b = select(v, _jd(v), vault_revision=3, bucket="high")
    assert a == b
    assert a.vault_revision == 3 and a.tailoring == "heuristic"
    assert a.projects[0].id == "proj-job-bot"            # FastAPI/Postgres/Redis backend
    assert [s.id for s in a.experience] == ["exp-acme-intern"]
    assert {"python", "fastapi", "postgresql"} <= set(a.keywords)


def test_select_changes_with_the_job(v):
    r = select(v, _jd(v, DS_JD, "Data Engineer", "Contoso"), vault_revision=1)
    assert r.projects[0].id == "proj-retail"


def test_select_reports_missing_must_haves(v):
    jd = _jd(v, "Requirements:\n- Kubernetes\n- Rust\n- Python", "Platform Engineer", "X")
    r = select(v, jd, vault_revision=1)
    assert "kubernetes" in r.missing_skills and "rust" in r.missing_skills
    assert "python" not in r.missing_skills


def test_select_respects_line_budget(v):
    r = select(v, _jd(v), vault_revision=1, line_budget=30)
    assert estimate_lines(v, r) <= 30 or sum(len(s.bullets) for s in r.projects) == 2 * len(r.projects)
    assert all(len(s.bullets) >= 2 for s in r.projects + r.experience)


def test_skills_put_jd_matches_first_and_keep_user_spelling(v):
    r = select(v, _jd(v), vault_revision=1)
    data = next(g for g in r.skills if g.category == "Data")
    assert data.items[0] == "PostgreSQL"          # user's spelling, JD-matched first


def test_recipe_is_small(v):
    r = build_recipe(v, _jd(v), vault_revision=1, bucket="high")
    assert r.size_bytes() < 4096


# --- truth validator -------------------------------------------------------------
SRC = "Cut p95 latency of the claim-status API from 900 ms to 350 ms by adding Redis caching and query indexes"
TECH = {"c#", "asp.net core", "sql server", "redis", "xunit"}


@pytest.mark.parametrize("rewrite,ok", [
    (SRC, True),
    ("Reduced p95 latency of the claim-status API from 900 ms to 350 ms with Redis caching and query indexes", True),
    ("Cut p95 latency of the claim-status API from 900 ms to 300 ms with Redis caching", False),   # changed number
    ("Cut p95 latency by 61% using Redis caching and query indexes", False),                      # derived number
    ("Cut p95 latency of the API from 900 ms to 350 ms using Kafka and Redis", False),             # new skill
    ("Cut p95 latency of the claim-status API from 900 ms to 350 ms using SQL Server indexes", True),  # owner tech
    ("Led a team of engineers to redesign the platform architecture for global scale", False),      # new content
    (SRC + " while mentoring junior engineers across multiple teams", False),                        # grew too much
])
def test_check_bullet(rewrite, ok):
    assert (check_bullet(SRC, rewrite, TECH, ()) is None) == ok


def test_metric_moved_from_another_bullet_is_rejected(v):
    # 60% is a real vault number (proctor review time) but not this bullet's.
    src = next(b.text for p in v.projects for b in p.bullets if b.id == "jb-3")
    assert "60%" in " ".join(extract_numbers(" ".join(b.text for p in v.projects for b in p.bullets)) + ["60%"])
    reason = check_bullet(src, src + " cutting review time by 60%", {"python"}, ())
    assert reason and "numbers" in reason


def test_validate_reverts_strips_and_drops(v):
    r = select(v, _jd(v), vault_revision=1)
    r.projects[0].bullets[0].text = "Built a Kubernetes operator handling 10M requests"   # invented
    r.projects.append(RSection(id="proj-ghost", bullets=[RBullet(id="x", text="fake")]))
    r.experience[0].bullets.append(RBullet(id="not-a-bullet"))
    r.skills.append(RSkillGroup(category="Cloud", items=["AWS", "postgres"]))
    r.keywords += ["kubernetes"]
    r.achievements += ["ach-ghost"]
    out = validate(v, r)
    assert out.projects[0].bullets[0].text is None                    # reverted to vault text
    assert "proj-ghost" in out.validator.dropped_ids
    assert "not-a-bullet" in out.validator.dropped_ids
    assert "AWS" in out.validator.stripped_skills and "kubernetes" in out.validator.stripped_skills
    cloud = next(g for g in out.skills if g.category == "Cloud")
    assert cloud.items == ["PostgreSQL"]                               # alias -> user's spelling
    assert "ach-ghost" not in out.achievements
    assert not out.validator.clean


def test_validate_removes_untraceable_sentences_only(v):
    r = select(v, _jd(v), vault_revision=1)
    r.cover_letter = ("I am applying for the role. I scaled systems to 5 million users. "
                      "I built 6 REST endpoints. I am an expert in Kubernetes.")
    r.answers = {"Why us?": "I like payments. I have 10 years of Go."}
    out = validate(v, r)
    assert out.cover_letter == "I am applying for the role. I built 6 REST endpoints."
    assert out.answers["Why us?"] == "I like payments."
    assert len(out.validator.removed_sentences) == 3


def test_employer_and_title_are_not_skill_claims(v):
    r = select(v, _jd(v, title="Java Developer", company="Oracle"), vault_revision=1)
    r.title, r.company = "Java Developer", "Oracle"
    r.cover_letter = "I am excited to apply for the Java Developer role at Oracle."
    assert validate(v, r).cover_letter == r.cover_letter


def test_validate_is_idempotent(v):
    r = build_recipe(v, _jd(v), vault_revision=1, bucket="high")
    assert validate(v, r) == r


# --- LLM tailoring with a scripted model ------------------------------------------
def _scripted(reply: dict) -> LLM:
    body = {"content": [{"type": "text", "text": json.dumps(reply)}]}
    return LLM(provider="anthropic", api_key="k", base_url="https://llm.test", cheap_model="c",
               strong_model="s", transport=httpx.MockTransport(lambda r: httpx.Response(200, json=body)))


def test_tailor_accepts_honest_rephrase_and_reverts_lies(v):
    reply = {
        "experience": [{"id": "exp-acme-intern", "bullets": [
            {"id": "acme-2", "text": "Reduced p95 latency of the claim-status API from 900 ms to 350 ms with Redis caching and query indexes"},
            {"id": "acme-1", "text": "Built 6 REST endpoints in ASP.NET Core serving 2 million requests/day"},
        ]}],
        "projects": [
            {"id": "proj-job-bot", "bullets": [{"id": "jb-2", "text": None}, {"id": "jb-1", "text": None}]},
            {"id": "proj-made-up", "bullets": [{"id": "zz", "text": "Invented"}]},
            {"id": "proj-netsec", "bullets": [{"id": "ns-3", "text": None}, {"id": "ns-1", "text": None}]},
        ],
        "skill_order": ["Python", "PostgreSQL", "FastAPI"],
    }
    jd = _jd(v)
    base = select(v, jd, vault_revision=1, bucket="high")
    out = validate(v, tailor(v, jd, base, _scripted(reply)))
    exp = {b.id: b.text for b in out.experience[0].bullets}
    assert exp["acme-2"].startswith("Reduced p95")          # honest rephrase kept
    assert exp["acme-1"] is None                            # inflated number reverted
    assert [s.id for s in out.projects] == ["proj-job-bot", "proj-netsec"]
    assert out.tailoring == "llm"
    assert out.skills[0].items[0] == "Python"


def test_tailor_falls_back_on_garbage(v):
    jd = _jd(v)
    base = select(v, jd, vault_revision=1, bucket="high")
    bad = LLM(provider="openai", api_key="k", base_url="https://llm.test", cheap_model="c", strong_model="s",
              transport=httpx.MockTransport(lambda r: httpx.Response(200, json={"choices": [{"message": {"content": "nope"}}]})))
    assert tailor(v, jd, base, bad) == base


def test_low_bucket_never_calls_llm(v):
    calls = []
    llm = LLM(provider="openai", api_key="k", base_url="https://llm.test", cheap_model="c", strong_model="s",
              transport=httpx.MockTransport(lambda r: calls.append(r) or httpx.Response(500)))
    build_recipe(v, _jd(v), vault_revision=1, bucket="low", llm=llm)
    assert calls == []


# --- cover letter / answers --------------------------------------------------------
def test_template_cover_letter_is_truthful(v):
    r = select(v, _jd(v), vault_revision=1)
    letter = template_cover_letter(v, r)
    assert "Northwind Payments" in letter and v.personal.full_name in letter
    assert validate(v, r.model_copy(update={"cover_letter": letter})).validator.removed_sentences == []


def test_standard_answers(v):
    a = standard_answers(v)
    assert a["Expected CTC"] == "12 LPA" and a["Notice period"] == "Immediate"
    assert a["Requires visa sponsorship"] == "No"


def test_cached_recipe_is_retargeted(v):
    first = build_recipe(v, _jd(v), vault_revision=1, bucket="high")
    jd2 = _jd(v, company="Contoso")
    second = build_recipe(v, jd2, vault_revision=1, bucket="medium", cached=first)
    assert second.tailoring == "cached" and second.company == "Contoso"
    assert "Contoso" in second.cover_letter and "Northwind" not in second.cover_letter
    assert [s.id for s in second.projects] == [s.id for s in first.projects]


def test_stale_cache_revision_ignored(v):
    first = build_recipe(v, _jd(v), vault_revision=1)
    assert build_recipe(v, _jd(v), vault_revision=2, cached=first).tailoring == "heuristic"


# --- one-page trimming ---------------------------------------------------------------
def test_drop_one_order(v):
    r = select(v, _jd(v), vault_revision=1)
    before = sum(len(s.bullets) for s in r.projects + r.experience)
    lowest = min((b.rank for s in r.projects + r.experience for b in s.bullets if len(s.bullets) > 2))
    assert _drop_one(r)
    after = [b.rank for s in r.projects + r.experience for b in s.bullets]
    assert len(after) == before - 1 and lowest not in after or after.count(lowest) < 2
    while _drop_one(r):
        pass
    assert len(r.projects) == 2 and all(len(s.bullets) == 2 for s in r.projects + r.experience)
    assert r.achievements == []


# --- rendering ----------------------------------------------------------------------
def test_latex_escape_all_specials():
    assert latex_escape(r"50% & $5 #1 a_b {x} ~ ^ \ <|>") == (
        r"50\% \& \$5 \#1 a\_b \{x\} \textasciitilde{} \textasciicircum{} \textbackslash{} "
        r"\textless{}\textbar{}\textgreater{}")
    assert latex_escape("₹12 “quoted” – dash") == "INR 12 ``quoted'' -- dash"


def test_latex_injection_is_neutralized(v):
    d = copy.deepcopy(EXAMPLE)
    d["projects"][0]["bullets"][0]["text"] = r"Pwned \input{/etc/passwd} \immediate\write18{rm -rf /}"
    v2 = Vault.model_validate(d)
    r = select(v2, _jd(v2), vault_revision=1)
    r.projects[0].bullets.insert(0, RBullet(id="jb-1", rank=9))
    tex = to_latex(build_view(v2, r))
    assert r"\input{/etc/passwd}" not in tex and r"\write18" not in tex
    assert r"\textbackslash{}input\{/etc/passwd\}" in tex


def test_typst_strings_are_inert():
    assert typst_string('#set page(fill: red) "x" \\') == '"#set page(fill: red) \\"x\\" \\\\"'


def test_text_and_typst_render(v):
    r = build_recipe(v, _jd(v), vault_revision=1)
    view = build_view(v, r)
    txt = to_text(view)
    assert txt.startswith("Sample Candidate\n") and "EXPERIENCE" in txt and "PROJECTS" in txt
    # fixed layout: summary, experience, skills, projects, education
    assert txt.index("EXPERIENCE") < txt.index("TECHNICAL SKILLS") < txt.index("PROJECTS") < txt.index("EDUCATION")
    assert "#" in to_typst(view)   # renders without template errors


def test_docx_twin_same_content_no_tables(v):
    from docx import Document
    r = build_recipe(v, _jd(v), vault_revision=1)
    view = build_view(v, r)
    doc = Document(io.BytesIO(to_docx(view)))
    assert doc.tables == []
    text = "\n".join(p.text for p in doc.paragraphs)
    for b in view.bullet_texts():
        assert b in text
    assert doc.paragraphs[0].text == "Sample Candidate"


def test_view_helpers():
    assert date_range("2026-07", None) == "Jul 2026 -- Present"
    assert date_range("2023-08", "2027-05", current_ok=False) == "Aug 2023 -- May 2027"
    assert display_url("https://www.github.com/x/") == "github.com/x"


def test_experience_years_merges_overlaps(v):
    from app.resume.validator import experience_years
    assert round(experience_years(v), 2) == round(4 / 12, 2)    # Jul-Oct 2026
    d = copy.deepcopy(EXAMPLE)
    d["experience"].append({**d["experience"][0], "id": "exp-2", "start": "2026-09", "end": "2026-12",
                            "bullets": [{"id": "e2-1", "text": "Overlapping role bullet text"}]})
    assert round(experience_years(Vault.model_validate(d)), 2) == round(6 / 12, 2)


def test_years_claim_capped_by_vault(v):
    r = select(v, _jd(v), vault_revision=1)
    r.answers = {"Experience?": "I have 0 years of full-time experience. I have 2+ years of Python."}
    assert validate(v, r).answers["Experience?"] == "I have 0 years of full-time experience."


def test_magnitude_words_are_numbers():
    assert extract_numbers("5 million users, 2 lakh orders, 1.5 crore") == ["5000000", "200000", "15000000"]
    assert extract_numbers("3 Mbps link") == ["3"]
