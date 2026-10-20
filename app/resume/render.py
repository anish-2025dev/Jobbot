"""Render a ResumeView to LaTeX source, Typst source, DOCX bytes or plain text."""
from __future__ import annotations

import io
import json
import re
from functools import lru_cache
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, StrictUndefined

from app.resume.model import ResumeView

TEMPLATES = Path(__file__).parent / "templates"

_LATEX_SPECIAL = {
    "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
    "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
    "<": r"\textless{}", ">": r"\textgreater{}", "|": r"\textbar{}",
}
_UNICODE_FIXES = {
    "₹": "INR ", "“": "``", "”": "''", "‘": "`", "’": "'", "–": "--", "—": "---",
    "…": "...", "•": "-", "→": "->", "​": "", "\xa0": " ",
}
_LATEX_RE = re.compile("|".join(re.escape(k) for k in _LATEX_SPECIAL))


def latex_escape(value: object) -> str:
    text = str(value)
    for k, rep in _UNICODE_FIXES.items():
        text = text.replace(k, rep)
    text = "".join(ch for ch in text if ch == "\n" or ord(ch) >= 32)
    # "--" is our own date separator; keep it as an en dash.
    return _LATEX_RE.sub(lambda m: _LATEX_SPECIAL[m.group()], text)


def latex_href(url: object) -> str:
    """Escape a link target for \\href{}: hyperref takes % and # escaped; anything that
    could close the group never gets here (model._safe_href)."""
    text = str(url)
    if any(ch in text for ch in "{}\\ \n"):
        raise ValueError("unsafe link target")
    return text.replace("%", r"\%").replace("#", r"\#")


def typst_string(value: object) -> str:
    text = "".join(ch for ch in str(value) if ch == "\n" or ord(ch) >= 32)
    text = text.replace(" -- ", " – ")
    return json.dumps(text, ensure_ascii=False)


@lru_cache
def _env(kind: str) -> Environment:
    env = Environment(
        loader=FileSystemLoader(str(TEMPLATES)),
        block_start_string="((*", block_end_string="*))",
        variable_start_string="(((", variable_end_string=")))",
        comment_start_string="((#", comment_end_string="#))",
        undefined=StrictUndefined, autoescape=False, keep_trailing_newline=True,
    )
    env.filters["e"] = latex_escape
    env.filters["href"] = latex_href
    env.filters["s"] = typst_string
    return env


def to_latex(view: ResumeView, paper: str = "a4paper") -> str:
    return _env("tex").get_template("resume.tex.j2").render(
        name=view.name, contact=view.contact, contact_links=_links(view), sections=view.sections, paper=paper)


def to_typst(view: ResumeView, paper: str = "a4paper") -> str:
    typst_paper = {"a4paper": "a4", "letterpaper": "us-letter"}.get(paper, "a4")
    return _env("typ").get_template("resume.typ.j2").render(
        name=view.name, contact=view.contact, contact_links=_links(view), sections=view.sections, paper=typst_paper)


def _links(view: ResumeView) -> list[tuple[str, str | None]]:
    return view.contact_links or [(c, None) for c in view.contact]


def to_text(view: ResumeView) -> str:
    """Plain text in reading order (what an ATS should extract)."""
    out = [view.name, " | ".join(view.contact), ""]
    for s in view.sections:
        out.append(s.title.upper())
        if s.paragraph:
            out.append(s.paragraph)
        for e in s.entries:
            head = e.heading + (f", {e.sub}" if e.sub else "")
            out.append(f"{head}  {e.right.replace('--', '-')}".rstrip())
            if e.second:
                out.append(e.second)
            out += [f"- {b}" for b in e.bullets]
        out += [f"{label}: {text}" for label, text in s.lines]
        out += [f"- {i}" for i in s.items]
        out.append("")
    return "\n".join(out).strip() + "\n"


def to_docx(view: ResumeView, paper: str = "a4paper") -> bytes:
    """DOCX twin for portals that reject PDF: same content, no tables."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document()
    sec = doc.sections[0]
    if paper == "a4paper":
        sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    else:
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.6))
    sec.top_margin = sec.bottom_margin = Inches(0.5)
    usable = sec.page_width - sec.left_margin - sec.right_margin

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    def para(text: str = "", bold: bool = False, size: float | None = None, align=None, italic=False):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        if text:
            run = p.add_run(text)
            run.bold, run.italic = bold, italic
            if size:
                run.font.size = Pt(size)
        return p

    def rule_below(p) -> None:
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, val in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "000000")):
            bottom.set(qn(k), val)
        bdr.append(bottom)
        pPr.append(bdr)

    para(view.name, bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(" | ".join(view.contact), align=WD_ALIGN_PARAGRAPH.CENTER)
    for s in view.sections:
        h = para(s.title.upper(), bold=True, size=11.5)
        h.paragraph_format.space_before = Pt(6)
        rule_below(h)
        if s.paragraph:
            para(s.paragraph)
        for e in s.entries:
            p = doc.add_paragraph()
            p.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
            p.add_run(e.heading).bold = True
            if e.sub:
                p.add_run(", ")
                p.add_run(e.sub).italic = True
            if e.right:
                p.add_run(f"\t{e.right.replace('--', '–')}")
            if e.second:
                para(e.second, italic=True)
            for b in e.bullets:
                doc.add_paragraph(b, style="List Bullet")
        for label, text in s.lines:
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(text)
        for it in s.items:
            doc.add_paragraph(it, style="List Bullet")
    buf = io.BytesIO()
    doc.core_properties.author = view.name
    doc.core_properties.title = f"{view.name} - Resume"
    doc.save(buf)
    return buf.getvalue()
